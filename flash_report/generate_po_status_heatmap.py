#!/usr/bin/env python3
"""
Generate a per-product PO-status heatmap workbook (one sheet per product group)
from the Flash Report KPI 2 datasource.

Layout (matches the shared "PO Status Sample.xlsx"):
  - One worksheet per product group.
  - Rows   = one per MPN / CatRef (the RSPL part), sorted by MPN.
  - Columns = that group's sites.
  - Each cell = the part's Action-Priority status at that site, color-coded.

Status / color scheme (per user spec, NOT the sample's wording):
  Action Priority -> cell text          -> color
  ---------------    ------------------    -----------------------
  (covered, OH>0) -> "On Hand"          -> green
  P5              -> "Recent PO, No PR" -> green
  P6              -> "PO & PR"          -> green
  P1              -> "No PO & No PR"    -> red
  P2              -> "Outdated PO, No PR" -> red
  P4              -> "PR, No PO"        -> red
  P3              -> "APM-only PO"      -> yellow
  Not in APM      -> "Not in APM"       -> yellow

Reuses the exact KPI 2 inclusion + Action Priority logic from
generate_flash_report.py so the heatmap and the report agree.

Reads a fresh pull of flash_report_all_queries.sql (flash_report_datasource_live.csv),
which already carries the split product labels (USP, USP B1, USP EU, URL Rev C/C2/D).
Each sheet's rows/columns come straight from that product's own rows.
"""

from datetime import date

import pandas as pd
import openpyxl
from openpyxl.styles import PatternFill, Font, Alignment, Border, Side
from openpyxl.utils import get_column_letter

SRC = "flash_report_datasource_live.csv"  # fresh pull of flash_report_all_queries.sql
OUT = "po_status_heatmap.xlsx"

# ---- colors (classic Excel traffic-light palette) -----------------------
GREEN_FILL = PatternFill("solid", fgColor="C6EFCE")
GREEN_FONT = Font(color="006100")
YELLOW_FILL = PatternFill("solid", fgColor="FFEB9C")
YELLOW_FONT = Font(color="9C5700")
RED_FILL = PatternFill("solid", fgColor="FFC7CE")
RED_FONT = Font(color="9C0006")
GREY_FILL = PatternFill("solid", fgColor="F2F2F2")
GREY_FONT = Font(color="808080")
HDR_FILL = PatternFill("solid", fgColor="232F3E")
HDR_FONT = Font(color="FFFFFF", bold=True)
THIN = Side(style="thin", color="D9D9D9")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)
CENTER = Alignment(horizontal="center", vertical="center")

# Each sheet maps to a product value in the datasource. The live datasource
# already carries the split product labels (USP B1, USP EU, URL Rev C/C2/D),
# so rows and columns come straight from each product's own data.
SHEETS = [
    ("URL Rev D",    "URL Rev D"),
    ("URL Rev C.2",  "URL Rev C2"),
    ("URL Rev C",    "URL Rev C"),
    ("USP",          "USP"),
    ("USP Bundle 1", "USP B1"),
    ("USP EU",       "USP EU"),
]

TODAY = pd.Timestamp(date.today())


# Descriptive cell labels (replace the P1..P6 codes). P3 uses "APM-only PO".
LBL_P1 = 'No PO & No PR'
LBL_P2 = 'Outdated PO, No PR'
LBL_P3 = 'APM-only PO'
LBL_P4 = 'PR, No PO'
LBL_P5 = 'Recent PO, No PR'
LBL_P6 = 'PO & PR'


def action_priority(hp_pr, hp_po, po_num, po_dt):
    """KPI 2 logic from generate_flash_report.py, returning descriptive labels."""
    if hp_pr == 'No' and hp_po == 'No':
        return LBL_P1
    if hp_pr == 'No' and hp_po == 'Yes':
        if str(po_num).strip().startswith('101'):
            return LBL_P3
        d = pd.to_datetime(po_dt, errors='coerce')
        if pd.notna(d) and (TODAY - d).days <= 30:
            return LBL_P5
        return LBL_P2
    if hp_pr == 'Yes' and hp_po == 'No':
        return LBL_P4
    if hp_pr == 'Yes' and hp_po == 'Yes':
        return LBL_P6
    return ''


def style_for(status):
    """Return (fill, font) for a cell status."""
    if status == 'On Hand':
        return GREEN_FILL, GREEN_FONT
    if status in (LBL_P5, LBL_P6):
        return GREEN_FILL, GREEN_FONT
    if status in (LBL_P1, LBL_P2, LBL_P4):
        return RED_FILL, RED_FONT
    if status in (LBL_P3, 'Not in APM'):
        return YELLOW_FILL, YELLOW_FONT
    return GREY_FILL, GREY_FONT  # not an RSPL part at this site


def build_cell_status(df):
    """
    For every (product, site, mpn, catalog_reference) group, compute the cell
    status using the same inclusion rule as KPI 2:

      - A group is "covered" (On Hand) if ANY of its APNs has OH > 0.
      - Otherwise the group is unavailable: compute Action Priority from the
        best (lowest-priority-number) APN row, i.e. the most actionable state.
      - If the group is not in APM (no APN found), status = "Not in APM".

    Returns a dict keyed (product, site, mpn, catref) -> status string.
    """
    oh = pd.to_numeric(df['site_oh_qty'], errors='coerce')
    df = df.assign(_oh=oh)

    grp_cols = ['product', 'site', 'mpn', 'catalog_reference']
    covered = df.groupby(grp_cols)['_oh'].transform(
        lambda x: (x.fillna(0) > 0).any())
    df = df.assign(_covered=covered)

    # Has PR / Has PO per row
    has_pr_num = pd.to_numeric(df.get('has_active_pr'), errors='coerce').fillna(0)
    df = df.assign(_has_pr=has_pr_num.apply(lambda v: 'Yes' if v == 1 else 'No'))
    po_num = df.get('po_number')
    df = df.assign(_has_po=po_num.apply(
        lambda v: 'Yes' if (pd.notna(v) and str(v).strip() not in ('', 'nan')) else 'No'))

    not_in_apm = df['mapping_status'] == 'No APN found'
    df = df.assign(_nia=not_in_apm)

    # Row-level action priority (only meaningful for in-APM, uncovered rows)
    df = df.assign(_ap=[action_priority(pr, po, pn, pd_)
                        for pr, po, pn, pd_ in
                        zip(df['_has_pr'], df['_has_po'], df['po_number'], df['po_date'])])

    prio_rank = {LBL_P1: 1, LBL_P2: 2, LBL_P3: 3, LBL_P4: 4, LBL_P5: 5, LBL_P6: 6, '': 9}

    status = {}
    for (product, site, mpn, catref), g in df.groupby(grp_cols, dropna=False):
        key = (product, site, str(mpn).strip(), str(catref).strip())
        if g['_covered'].iloc[0]:
            status[key] = 'On Hand'
        elif g['_nia'].all():
            status[key] = 'Not in APM'
        else:
            # pick the most actionable (lowest rank) non-blank priority
            ranked = g[g['_ap'] != ''].copy()
            if ranked.empty:
                status[key] = 'Not in APM'
            else:
                ranked['_r'] = ranked['_ap'].map(prio_rank).fillna(9)
                status[key] = ranked.sort_values('_r').iloc[0]['_ap']
    return status


def product_pairs(df):
    """Distinct MPN/CatRef pairs for a product's rows, sorted by MPN.

    MPN/CRN are stripped of stray whitespace (a few source rows carry a
    trailing-space artifact, e.g. 'SC25EUM-5 ').
    """
    sub = df[['mpn', 'catalog_reference']].copy()
    sub['mpn'] = sub['mpn'].fillna('').astype(str).str.strip()
    sub['catalog_reference'] = sub['catalog_reference'].fillna('').astype(str).str.strip()
    pairs = (sub.drop_duplicates().sort_values(['mpn', 'catalog_reference']))
    return list(pairs.itertuples(index=False, name=None))


def write_sheet(wb, title, product, pairs, status_map, sites):
    ws = wb.create_sheet(title=title)
    # header
    headers = ['MPN', 'CatRef'] + sites
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=1, column=c, value=h)
        cell.fill = HDR_FILL
        cell.font = HDR_FONT
        cell.alignment = CENTER
        cell.border = BORDER

    for r, (mpn, catref) in enumerate(pairs, start=2):
        ws.cell(row=r, column=1, value=mpn).border = BORDER
        ws.cell(row=r, column=2, value=catref).border = BORDER
        for ci, site in enumerate(sites, start=3):
            # Status from the datasource for this product/site/part. A blank
            # means the part isn't on this site's RSPL -> leave grey/blank.
            st = status_map.get((product, site, str(mpn), str(catref)), '')
            cell = ws.cell(row=r, column=ci, value=st)
            fill, font = style_for(st)
            cell.fill = fill
            cell.font = font
            cell.alignment = CENTER
            cell.border = BORDER

    # column widths + freeze panes
    ws.column_dimensions['A'].width = 24
    ws.column_dimensions['B'].width = 14
    for ci in range(3, len(headers) + 1):
        ws.column_dimensions[get_column_letter(ci)].width = 18
    ws.freeze_panes = "C2"
    return len(pairs)


def build_legend(wb):
    ws = wb.create_sheet(title="Legend", index=0)
    ws.column_dimensions['A'].width = 16
    ws.column_dimensions['B'].width = 70
    ws['A1'] = 'PO Status Heatmap'
    ws['A1'].font = Font(bold=True, size=14)
    ws['A3'] = 'Status'
    ws['B3'] = 'Meaning'
    for c in ('A3', 'B3'):
        ws[c].fill = HDR_FILL
        ws[c].font = HDR_FONT
    rows = [
        ('On Hand', 'green', 'Part is stocked at this site (on-hand qty > 0).'),
        (LBL_P5, 'green', 'Recent PO (<=30 days), no PR - no action required.'),
        (LBL_P6, 'green', 'Has both PO and PR - no action required.'),
        (LBL_P1, 'red', 'No PO and no PR - nothing in motion. Act first.'),
        (LBL_P2, 'red', 'Outdated PO (>30 days), no PR.'),
        (LBL_P4, 'red', 'Has PR, no PO - PO pending.'),
        (LBL_P3, 'yellow', 'APM-only PO (PO # starts 101), no PR - needs data validation.'),
        ('Not in APM', 'yellow', 'RSPL part has no APN mapping in APM.'),
        ('(blank/grey)', 'grey', 'Part is not on this site\'s RSPL.'),
    ]
    for i, (label, color, meaning) in enumerate(rows, start=4):
        a = ws.cell(row=i, column=1, value=label)
        b = ws.cell(row=i, column=2, value=meaning)
        if color == 'green':
            a.fill, a.font = GREEN_FILL, GREEN_FONT
        elif color == 'red':
            a.fill, a.font = RED_FILL, RED_FONT
        elif color == 'yellow':
            a.fill, a.font = YELLOW_FILL, YELLOW_FONT
        else:
            a.fill, a.font = GREY_FILL, GREY_FONT
        a.alignment = CENTER


def main():
    df = pd.read_csv(SRC, dtype=str, keep_default_na=False, na_filter=False)
    status_map = build_cell_status(df)

    wb = openpyxl.Workbook()
    # drop the default sheet; legend created with index=0
    wb.remove(wb.active)
    build_legend(wb)

    for title, product in SHEETS:
        sub = df[df['product'] == product].copy()
        sites = sorted(sub['site'].dropna().unique().tolist())
        pairs = product_pairs(sub)
        n = write_sheet(wb, title, product, pairs, status_map, sites)
        print(f"  {title}: {n} MPNs x {len(sites)} sites")

    wb.save(OUT)
    print(f"Saved {OUT}")


if __name__ == "__main__":
    main()
